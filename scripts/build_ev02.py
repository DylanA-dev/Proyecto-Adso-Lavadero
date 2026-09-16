# -*- coding: utf-8 -*-
"""
Construye la evidencia GA6-220501096-AA3-EV02 (mockups de navegación - Lavadero).

Siempre parte de la guia original (1GUIA_PARA_SENA .docx) para que sea
idempotente: se puede ejecutar varias veces sin duplicar contenido.
Inserta las capturas de la carpeta prototipo/capturas si existen.
"""
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ORIGINAL = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\1GUIA_PARA_SENA .docx"
OUTPUT = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\GA6-220501096-AA3-EV02.docx"
CAPTURAS = Path(r"C:\Users\dylan\Documents\OpenCode\ADSO\proyectos\lavadero\prototipo\capturas")

CODIGO = "GA6-220501096-AA3-EV02"
TITULO_EVIDENCIA = "Diseño del sitio web y/o móviles utilizando sus componentes y tecnologías respectivas"
FECHA = "Agosto 2026"

# ------------------------- Utilidades -------------------------

def add_paragraph_after(anchor, text="", style=None, align=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = style
    if align is not None:
        para.alignment = align
    return para


def add_picture_after(anchor, image_path, width_cm=15.0, caption=None):
    para = add_paragraph_after(anchor, "", "Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        cap = add_paragraph_after(para, caption, "Normal",
                                  align=WD_ALIGN_PARAGRAPH.CENTER)
        # Cursiva en el caption
        for r in cap.runs:
            r.italic = True
        return cap
    return para


def normalize(text):
    import unicodedata
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().upper()


def find_paragraph(doc, target):
    """Busca un parrafo cuyo texto normalizado contenga el target."""
    t = normalize(target)
    for p in doc.paragraphs:
        if t in normalize(p.text):
            return p
    return None


# ------------------------- 1. Copia fresca desde la guia original -------------------------
shutil.copy2(ORIGINAL, OUTPUT)
doc = Document(OUTPUT)

# ------------------------- 2. Portada -------------------------
title_p = find_paragraph(doc, "Evidencia de conocimiento")
if title_p:
    # Limpiar runs existentes y poner el codigo de la actividad como titulo
    for r in list(title_p.runs):
        r._r.getparent().remove(r._r)
    title_p.add_run(CODIGO)
    anchor = add_paragraph_after(title_p, TITULO_EVIDENCIA, "Normal",
                                 align=WD_ALIGN_PARAGRAPH.CENTER)

fecha_p = find_paragraph(doc, "MES")
if fecha_p:
    for r in list(fecha_p.runs):
        r._r.getparent().remove(r._r)
    fecha_p.add_run(FECHA)

# ------------------------- 3. Introduccion -------------------------
intro_h = find_paragraph(doc, "INTRODUCCION")
intro_text = (
    "En el marco de la formación como tecnólogo en Análisis y Desarrollo de Software, "
    "el diseño de interfaces y la elaboración de prototipos de navegación constituyen "
    "una etapa fundamental dentro del ciclo de vida del desarrollo de aplicaciones, ya "
    "que permiten visualizar de manera anticipada la estructura, el flujo y la "
    "usabilidad del sistema antes de su implementación. La presente evidencia tiene "
    "como propósito dar a conocer los mockups de navegación diseñados para el sistema "
    "de gestión de un lavadero de vehículos, denominado AquaLavado, dando cumplimiento "
    "a los requerimientos funcionales planteados para el software."
)
intro_text2 = (
    "Para su elaboración se emplearon las herramientas de prototipado vistas en los "
    "componentes formativos, representando las pantallas principales del aplicativo "
    "web, entre las cuales se destacan el inicio de sesión, el panel principal, el "
    "registro de lavados, la gestión de clientes, el control de procesos y la "
    "generación de reportes. Cada una de estas interfaces se describe a continuación, "
    "evidenciando su funcionalidad, los elementos que la componen y la navegación "
    "que permite demostrar el funcionamiento del software."
)
if intro_h:
    anchor = add_paragraph_after(intro_h, intro_text, "Normal")
    anchor = add_paragraph_after(anchor, intro_text2, "Normal")

# ------------------------- 4. Desarrollo de la actividad -------------------------
dev_h = find_paragraph(doc, "DESARROLLO DE LA ACTIVIDAD")

# 4.1 Objetivo
anchor = add_paragraph_after(dev_h, "Objetivo", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Desarrollar los mockups de navegación de la aplicación web del sistema de gestión "
    "de lavadero de vehículos, con el fin de visualizar el flujo de las interfaces y "
    "demostrar la funcionalidad del software según sus requerimientos.", "Normal")
for obj in [
    "Representar gráficamente las pantallas principales del sistema: inicio de sesión, panel principal, registro de lavados, gestión de clientes, control de procesos y reportes.",
    "Establecer la navegación entre las diferentes vistas del aplicativo mediante enlaces funcionales dentro del prototipo.",
    "Describir cada interfaz y los elementos que la componen, garantizando una experiencia de usuario clara e intuitiva.",
]:
    anchor = add_paragraph_after(anchor, "• " + obj, "Normal")

# 4.2 Descripcion del software
anchor = add_paragraph_after(anchor, "Descripción del software", "Heading 2")
anchor = add_paragraph_after(anchor,
    "El sistema AquaLavado es una aplicación web de gestión para lavaderos de "
    "vehículos que permite administrar el registro de lavados, la información de los "
    "clientes, el control de los procesos de lavado en tiempo real y la generación de "
    "reportes e indicadores de la operación. El prototipo presentado a continuación "
    "muestra la navegación entre sus módulos principales y la funcionalidad de cada "
    "interfaz según las características del software a desarrollar.", "Normal")

# 4.3 Mockups: pantallazo + descripcion de cada interfaz
interfaces = [
    ("1_login.png", "Interfaz 1. Inicio de sesión (Login)",
     "Esta es la primera interfaz del aplicativo. Permite al usuario autenticarse en "
     "el sistema mediante un formulario que solicita usuario y contraseña. Al "
     "presionar el botón \u201cIniciar sesión\u201d se valida el acceso y se redirige "
     "al panel principal. Es la puerta de entrada que garantiza que únicamente el "
     "personal autorizado acceda a la información del lavadero."),
    ("2_dashboard.png", "Interfaz 2. Panel principal (Dashboard)",
     "Muestra un resumen general de la operación del lavadero: cantidad de lavados "
     "del día, ingresos generados, clientes registrados y lavados en proceso. "
     "Además, presenta un gráfico de lavados por tipo de servicio y accesos rápidos "
     "a las secciones de registro de lavados y control de procesos, lo que facilita "
     "la navegación del usuario desde el módulo de inicio."),
    ("3_lavados.png", "Interfaz 3. Registro de lavados",
     "Permite crear un nuevo servicio de lavado seleccionando el cliente, el "
     "vehículo, el tipo de lavado (básico, completo, premium o detallado) y la "
     "fecha. Incluye una tabla con los lavados recientes y su estado (pendiente, en "
     "cola, en proceso o finalizado), permitiendo hacer seguimiento de los servicios "
     "registrados en el sistema."),
    ("4_clientes.png", "Interfaz 4. Gestión de clientes",
     "Facilita el registro, la búsqueda y la consulta de los clientes del lavadero. "
     "Muestra un listado con los datos principales de cada cliente (nombre, cédula, "
     "teléfono, vehículo y último lavado) y permite registrar nuevos clientes "
     "mediante un formulario, manteniendo actualizada la base de datos de la "
     "operación."),
    ("5_procesos.png", "Interfaz 5. Control de procesos",
     "Presenta un tablero tipo kanban con las etapas del proceso de lavado: en cola, "
     "en proceso y finalizados. Cada tarjeta contiene el vehículo, el cliente, el "
     "operario asignado y el avance del lavado, lo que permite realizar el "
     "seguimiento de la operación en tiempo real y controlar la productividad del "
     "personal."),
    ("6_reportes.png", "Interfaz 6. Reportes y estadísticas",
     "Consolida los indicadores de la operación: ingresos del periodo, total de "
     "lavados, ticket promedio y clientes recurrentes. Incluye gráficos de ingresos "
     "por tipo de lavado y una tabla de lavados por día de la semana, información "
     "útil para la toma de decisiones y la mejora continua del servicio."),
]

for i, (img, titulo, desc) in enumerate(interfaces, start=1):
    anchor = add_paragraph_after(anchor, titulo, "Heading 2")
    img_path = CAPTURAS / img
    if img_path.exists():
        nombre_interfaz = titulo.split(". ", 1)[1].lower()
        anchor = add_picture_after(
            anchor, img_path, width_cm=14.5,
            caption=f"Figura {i}. Mockup de la interfaz {nombre_interfaz} — prototipo navegable.")
    else:
        anchor = add_paragraph_after(
            anchor, f"[Pantallazo pendiente: {img} — captúralo desde {img.replace('.png', '.html')}]",
            "Normal")
    anchor = add_paragraph_after(anchor, desc, "Normal")

# ------------------------- 5. Conclusion -------------------------
concl_h = find_paragraph(doc, "CONCLUSION")
concl_text = (
    "El desarrollo de la presente evidencia permitió aplicar los conocimientos "
    "adquiridos en los componentes formativos relacionados con el prototipado de "
    "interfaces, materializando los requerimientos del sistema de gestión de lavadero "
    "de vehículos en mockups de navegación funcionales. A través de este ejercicio se "
    "comprendió la importancia de planificar visualmente la estructura de una "
    "aplicación antes de su implementación, ya que esto facilita la detección de "
    "posibles errores de usabilidad, optimiza el tiempo de desarrollo y alinea las "
    "expectativas del cliente con el producto final."
)
concl_text2 = (
    "Así mismo, se fortalecieron competencias relacionadas con el diseño de "
    "experiencia de usuario, la organización de los módulos del software y la "
    "comunicación efectiva de las funcionalidades a través de cada interfaz. Se "
    "concluye que el prototipo cumple con los requerimientos establecidos y sirve "
    "como base sólida para la etapa de construcción e implementación del sistema "
    "AquaLavado."
)
if concl_h:
    anchor = add_paragraph_after(concl_h, concl_text, "Normal")
    anchor = add_paragraph_after(anchor, concl_text2, "Normal")

doc.save(OUTPUT)
print(f"OK: documento generado en {OUTPUT}")
missing = [img for img, _, _ in interfaces if not (CAPTURAS / img).exists()]
if missing:
    print("Faltan capturas:", ", ".join(missing))
else:
    print("Todas las capturas fueron insertadas.")

