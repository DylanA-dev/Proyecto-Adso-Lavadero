# -*- coding: utf-8 -*-
"""
Construye la evidencia GA6-220501096-AA4-EV02 (documento de componentes frontend - Lavadero).

Siempre parte de la guia original (1GUIA_PARA_SENA .docx) para que sea
idempotente: se puede ejecutar varias veces sin duplicar contenido.
Inserta capturas de la interfaz real (interfaz/capturas) si existen.
"""
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ORIGINAL = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\1GUIA_PARA_SENA .docx"
OUTPUT = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\GA6-220501096-AA4-EV02.docx"
CAPTURAS = Path(r"C:\Users\dylan\Documents\OpenCode\ADSO\proyectos\lavadero\interfaz\capturas")

CODIGO = "GA6-220501096-AA4-EV02"
TITULO_EVIDENCIA = "Establecer los componentes front-end de la aplicación web"
FECHA = "Agosto 2026"

# ------------------------- Utilidades -------------------------

def add_paragraph_after(anchor, text="", style=None, align=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = style
    if align is not None:
        para.alignment = align
    return para


def add_picture_after(anchor, image_path, width_cm=15.5, caption=None):
    para = add_paragraph_after(anchor, "", "Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        cap = add_paragraph_after(para, caption, "Normal",
                                  align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in cap.runs:
            r.italic = True
        return cap
    return para


def add_table_after(anchor, rows, cols, width_cols=None):
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)
    if width_cols:
        for row in tbl.rows:
            for idx, w in enumerate(width_cols):
                row.cells[idx].width = Cm(w)
    anchor._p.addnext(tbl._tbl)
    return tbl


def fill_table(tbl, data):
    for r, fila in enumerate(data):
        for c, val in enumerate(fila):
            cell = tbl.rows[r].cells[c]
            cell.text = val
            if r == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def anchor_after_table(tbl):
    new_p = OxmlElement("w:p")
    tbl._tbl.addnext(new_p)
    return Paragraph(new_p, tbl._parent)


def add_code_paragraph(anchor, code, style="Normal", align=None):
    para = add_paragraph_after(anchor, "", style, align)
    run = para.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Cm(0.40)
    return para


def add_code_block(anchor, code):
    for linea in code.split("\n"):
        anchor = add_code_paragraph(anchor, linea)
    return anchor


def normalize(text):
    import unicodedata
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().upper()


def find_paragraph(doc, target):
    t = normalize(target)
    for p in doc.paragraphs:
        if t in normalize(p.text):
            return p
    return None


# ------------------------- 1. Copia fresca desde la guia original -------------------------
shutil.copy2(ORIGINAL, OUTPUT)
doc = Document(OUTPUT)

# ------------------------- 2. Portada -------------------------
title_p = find_paragraph(doc, "Evidencia de conocimiento")
if title_p:
    for r in list(title_p.runs):
        r._r.getparent().remove(r._r)
    title_p.add_run(CODIGO)
    anchor = add_paragraph_after(title_p, TITULO_EVIDENCIA, "Normal",
                                 align=WD_ALIGN_PARAGRAPH.CENTER)

fecha_p = find_paragraph(doc, "MES")
if fecha_p:
    for r in list(fecha_p.runs):
        r._r.getparent().remove(r._r)
    fecha_p.add_run(FECHA)

# ------------------------- 3. Introduccion -------------------------
intro_h = find_paragraph(doc, "INTRODUCCION")
intro_text = (
    "En el desarrollo de aplicaciones web, los componentes frontend definen la "
    "estructura, el contenido y la presentación de la interfaz con la que interactúa "
    "el usuario, por lo que su correcta definición resulta indispensable para "
    "garantizar que el software cumpla con los requerimientos funcionales planteados. "
    "La presente evidencia tiene como propósito establecer los componentes frontend "
    "que se utilizarán en el desarrollo de la aplicación web del sistema de gestión de "
    "lavadero de vehículos AquaLavado, especificando los elementos HTML que "
    "conformarán cada una de las interfaces del software."
)
intro_text2 = (
    "Para dar cumplimiento a la actividad se definió la estructura de componentes de la "
    "aplicación, partiendo de las características del software a desarrollar: un "
    "sistema que permite gestionar el registro de lavados, la información de los "
    "clientes, el control de los procesos de lavado y la generación de reportes. "
    "En el presente documento se especifican los elementos HTML a utilizar en el "
    "diseño frontend, la organización de los archivos del proyecto, los estilos CSS "
    "y los componentes JavaScript, así como la relación de cada componente con las "
    "pantallas del sistema."
)
if intro_h:
    anchor = add_paragraph_after(intro_h, intro_text, "Normal")
    anchor = add_paragraph_after(anchor, intro_text2, "Normal")

# ------------------------- 4. Desarrollo de la actividad -------------------------
dev_h = find_paragraph(doc, "DESARROLLO DE LA ACTIVIDAD")

# 4.1 Objetivo
anchor = add_paragraph_after(dev_h, "Objetivo", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Establecer la estructura de componentes frontend que se van a utilizar en el "
    "desarrollo de la aplicación web del sistema de gestión de lavadero de vehículos "
    "AquaLavado, definiendo los elementos HTML necesarios para cumplir con las "
    "necesidades del software a desarrollar.", "Normal")
for obj in [
    "Definir la estructura de archivos y la organización de los componentes frontend de la aplicación.",
    "Especificar los elementos HTML que se utilizarán en cada una de las interfaces del sistema.",
    "Describir los componentes de estilos CSS y las funcionalidades JavaScript aplicadas en el frontend.",
    "Relacionar los componentes definidos con las necesidades funcionales del software AquaLavado.",
]:
    anchor = add_paragraph_after(anchor, "• " + obj, "Normal")

# 4.2 Descripcion del software
anchor = add_paragraph_after(anchor, "Descripción del software", "Heading 2")
anchor = add_paragraph_after(anchor,
    "AquaLavado es una aplicación web de gestión para lavaderos de vehículos que "
    "permite administrar el registro de lavados, la información de los clientes, el "
    "control de los procesos de lavado en tiempo real y la generación de reportes e "
    "indicadores de la operación. El software se desarrolla con tecnologías frontend "
    "web (HTML5, CSS3 y JavaScript) y está conformado por los siguientes módulos: "
    "inicio de sesión, panel principal, registro de lavados, gestión de clientes, "
    "control de procesos y reportes.", "Normal")

# 4.3 Estructura de componentes / archivos
anchor = add_paragraph_after(anchor, "Estructura de componentes del frontend", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La aplicación se organiza en una estructura de archivos que separa el contenido "
    "HTML, los estilos CSS y la lógica JavaScript, facilitando su mantenimiento y "
    "escalabilidad. La estructura de componentes definida es la siguiente:", "Normal")
anchor = add_code_block(anchor,
    "interfaz/\n"
    "├── index.html          → Página de inicio de sesión\n"
    "├── app.html            → Aplicación principal (panel, lavados, clientes, procesos, reportes)\n"
    "├── css/\n"
    "│   └── estilo.css      → Hoja de estilos CSS del sistema\n"
    "└── js/\n"
    "    ├── datos.js        → Datos iniciales y configuración\n"
    "    └── app.js          → Lógica de la aplicación y componentes funcionales")

# 4.4 Elementos HTML a utilizar
anchor = add_paragraph_after(anchor, "Definición de elementos HTML a utilizar", "Heading 2")
anchor = add_paragraph_after(anchor,
    "De acuerdo con las características del software, a continuación se definen los "
    "elementos HTML que se van a utilizar en el diseño frontend de la aplicación, "
    "indicando su etiqueta, su propósito y el módulo o componente en el que se "
    "emplea:", "Normal")
tabla = add_table_after(anchor, 17, 3, width_cols=[3.5, 5.5, 6.5])
fill_table(tabla, [
    ("Elemento HTML", "Descripción", "Uso en el sistema AquaLavado"),
    ("<!DOCTYPE html>", "Declaración del tipo de documento HTML5.", "Todas las páginas del sistema."),
    ("<html>", "Elemento raíz del documento HTML.", "Todas las páginas del sistema."),
    ("<head>", "Cabecera con metadatos, título y enlaces a recursos.", "Todas las páginas (metadatos, css, js)."),
    ("<body>", "Cuerpo que contiene el contenido visible.", "Todas las páginas del sistema."),
    ("<header>", "Encabezado de una sección o de la página.", "Encabezado de la aplicación con la identidad del sistema."),
    ("<nav>", "Contenedor del menú de navegación.", "Menú lateral del panel de la aplicación."),
    ("<main>", "Contenido principal del documento.", "Contenedor de las vistas del sistema."),
    ("<section>", "Sección temática de la página.", "Módulos: panel, lavados, clientes, procesos y reportes."),
    ("<form>", "Formulario para capturar datos del usuario.", "Inicio de sesión, registro de lavados y clientes."),
    ("<input>", "Campo de entrada de datos (texto, número, fecha, contraseña).", "Usuario/contraseña, datos del cliente y del lavado."),
    ("<select>", "Lista desplegable de opciones.", "Tipo de servicio, cliente, operario y vehículo."),
    ("<textarea>", "Área de texto multilínea.", "Observaciones del lavado (opcional)."),
    ("<button>", "Botón de acción.", "Iniciar sesión, guardar lavado/cliente, avanzar proceso."),
    ("<table>", "Tabla de datos organizada en filas y columnas.", "Listado de lavados, clientes y reportes."),
    ("<img>", "Imagen en la interfaz.", "Logotipo del sistema y avatares de clientes."),
    ("<div>", "Contenedor genérico para agrupar y maquetar.", "Tarjetas de indicadores, kanban, paneles."),
])
anchor = anchor_after_table(tabla)

# 4.5 Componentes CSS
anchor = add_paragraph_after(anchor, "Componentes de estilos CSS", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Para cumplir con las necesidades del software se definen los siguientes "
    "componentes de estilos CSS, organizados en la hoja estilo.css:", "Normal")
tabla2 = add_table_after(anchor, 8, 3, width_cols=[4.0, 5.0, 6.5])
fill_table(tabla2, [
    ("Componente CSS", "Descripción", "Uso en el sistema"),
    ("Variables CSS (:root)", "Centralizan colores, fuentes y medidas.", "Paleta de colores del sistema (azul #0e3a7d, cian #22d3ee, etc.)."),
    ("Barra lateral (sidebar)", "Menú lateral fijo con las opciones del sistema.", "Navegación principal entre los módulos."),
    ("Tarjetas (cards)", "Contenedores con sombra y bordes redondeados.", "Indicadores del panel y tarjetas del kanban."),
    ("Tablas estilizadas", "Tablas con encabezado, filas alternadas y hover.", "Listados de lavados, clientes y reportes."),
    ("Formularios", "Campos, etiquetas y botones con estilos propios.", "Login, registro de lavados y de clientes."),
    ("Etiquetas de estado (badges)", "Indicadores de color por estado.", "Estados de lavado: en cola, en proceso, finalizado."),
    ("Media queries", "Adaptación del diseño a distintos tamaños.", "Diseño responsivo para móviles y tabletas."),
])
anchor = anchor_after_table(tabla2)

# 4.6 Componentes JavaScript
anchor = add_paragraph_after(anchor, "Componentes JavaScript", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La lógica del frontend se implementa en JavaScript para brindar funcionalidad a "
    "la interfaz. Los principales componentes funcionales son:", "Normal")
for js in [
    "Validación de sesión: controla el acceso al sistema mediante la comprobación de credenciales.",
    "Navegación por secciones: muestra u oculta las vistas del sistema sin recargar la página.",
    "Gestión de formularios: captura, valida y guarda los datos de lavados y clientes.",
    "Renderizado dinámico: genera las tablas, tarjetas y el tablero kanban a partir de los datos.",
    "Tablero kanban: permite avanzar el estado de cada lavado (en cola, en proceso, finalizado).",
    "Búsqueda en tiempo real: filtra clientes y lavados por nombre, cédula o vehículo.",
    "Reportes e indicadores: calcula ingresos, total de lavados y ticket promedio.",
    "Persistencia de datos: guarda la información en localStorage del navegador.",
]:
    anchor = add_paragraph_after(anchor, "• " + js, "Normal")

# 4.7 Componentes por modulo (relacion con pantallas)
anchor = add_paragraph_after(anchor, "Componentes HTML por módulo del sistema", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La siguiente tabla relaciona cada módulo de la aplicación con los elementos HTML "
    "y componentes que se utilizan para cumplir con sus requerimientos:", "Normal")
tabla3 = add_table_after(anchor, 7, 3, width_cols=[3.8, 5.7, 6.0])
fill_table(tabla3, [
    ("Módulo", "Requerimiento funcional", "Elementos HTML y componentes utilizados"),
    ("Inicio de sesión", "Permitir el acceso al sistema.", "<form>, <input>, <button>, <img>, <header>."),
    ("Panel principal", "Mostrar indicadores de la operación.", "<main>, <section>, <div> (tarjetas), <table>."),
    ("Registro de lavados", "Registrar un nuevo servicio de lavado.", "<form>, <select>, <input>, <button>, <table>."),
    ("Gestión de clientes", "Registrar y buscar clientes.", "<form>, <input>, <button>, <table>."),
    ("Control de procesos", "Seguimiento del estado de los lavados.", "<section>, <div> (kanban), <button>."),
    ("Reportes", "Generar estadísticas e indicadores.", "<main>, <section>, <table>, <div> (tarjetas)."),
])
anchor = anchor_after_table(tabla3)

# 4.8 Capturas de la interfaz implementada
anchor = add_paragraph_after(anchor, "Interfaz implementada con los componentes", "Heading 2")
anchor = add_paragraph_after(anchor,
    "A continuación se presentan pantallazos de la interfaz gráfica implementada, "
    "donde se evidencia la aplicación de los componentes frontend definidos en este "
    "documento:", "Normal")

interfaces = [
    ("1_login.png", "Vista 1. Inicio de sesión",
     "Muestra los componentes <form>, <input> y <button> utilizados para validar el "
     "acceso al sistema, junto con los estilos CSS del fondo y la tarjeta de inicio "
     "de sesión."),
    ("2_dashboard.png", "Vista 2. Panel principal",
     "Aplica <main> y <section> para organizar el contenido, tarjetas <div> con los "
     "indicadores y <nav> para el menú lateral del sistema."),
    ("3_lavados.png", "Vista 3. Registro de lavados",
     "Usa <form>, <select> e <input> para capturar el lavado y <table> para mostrar "
     "el listado de lavados recientes."),
    ("4_clientes.png", "Vista 4. Gestión de clientes",
     "Emplea <form> e <input> para registrar clientes y <table> con el listado, "
     "incluyendo el buscador en tiempo real."),
    ("5_procesos.png", "Vista 5. Control de procesos",
     "Implementa el tablero kanban con <section> y <div> por cada estado, con el "
     "botón <button> para avanzar los lavados."),
    ("6_reportes.png", "Vista 6. Reportes y estadísticas",
     "Organiza los indicadores en tarjetas <div> y los reportes en <table> para "
     "presentar las estadísticas de la operación."),
]

for i, (img, titulo, desc) in enumerate(interfaces, start=1):
    anchor = add_paragraph_after(anchor, titulo, "Heading 2")
    img_path = CAPTURAS / img
    if img_path.exists():
        anchor = add_picture_after(anchor, img_path, width_cm=15.0)
    else:
        anchor = add_paragraph_after(anchor, f"[Pantallazo pendiente: {img}]", "Normal")
    anchor = add_paragraph_after(anchor, desc, "Normal")

# 4.9 Archivos entregados
anchor = add_paragraph_after(anchor, "Documento entregado", "Heading 2")
anchor = add_paragraph_after(anchor,
    "De acuerdo con los lineamientos de la evidencia, el producto a entregar es el "
    "presente documento de componentes frontend, el cual se empaca en una carpeta "
    "comprimida en formato ZIP.", "Normal")

# ------------------------- 5. Conclusion -------------------------
concl_h = find_paragraph(doc, "CONCLUSION")
concl_text = (
    "El desarrollo de la presente evidencia permitió establecer la estructura de "
    "componentes frontend de la aplicación web AquaLavado, definiendo los elementos "
    "HTML, los estilos CSS y las funcionalidades JavaScript necesarios para cumplir "
    "con las características y necesidades del software a desarrollar. Se comprendió "
    "la importancia de especificar de forma clara cada componente antes de su "
    "implementación, ya que esto facilita el desarrollo, el mantenimiento y la "
    "escalabilidad de la aplicación."
)
concl_text2 = (
    "Así mismo, se fortalecieron competencias en la selección y organización de "
    "elementos HTML para el diseño frontend, relacionando cada componente con el "
    "módulo del sistema en el que se utiliza. Se concluye que los componentes "
    "definidos en este documento responden a los requerimientos funcionales de "
    "AquaLavado y constituyen la base técnica para la construcción del frontend "
    "de la aplicación web."
)
if concl_h:
    anchor = add_paragraph_after(concl_h, concl_text, "Normal")
    anchor = add_paragraph_after(anchor, concl_text2, "Normal")

doc.save(OUTPUT)
print(f"OK: documento generado en {OUTPUT}")
missing = [img for img, _, _ in interfaces if not (CAPTURAS / img).exists()]
if missing:
    print("Faltan capturas:", ", ".join(missing))
else:
    print("Todas las capturas fueron insertadas.")


