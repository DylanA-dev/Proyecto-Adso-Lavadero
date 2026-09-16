# -*- coding: utf-8 -*-
"""
Construye la evidencia GA6-220501096-AA3-EV03 (interfaces graficas - Lavadero).

Siempre parte de la guia original (1GUIA_PARA_SENA .docx) para que sea
idempotente. Inserta las capturas de la interfaz grafica real.
"""
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ORIGINAL = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\1GUIA_PARA_SENA .docx"
OUTPUT = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\GA6-220501096-AA3-EV03.docx"
CAPTURAS = Path(r"C:\Users\dylan\Documents\OpenCode\ADSO\proyectos\lavadero\interfaz\capturas")

CODIGO = "GA6-220501096-AA3-EV03"
TITULO_EVIDENCIA = "Interfaces gráficas según requerimientos del proyecto"
FECHA = "Agosto 2026"

# ------------------------- Utilidades -------------------------

def add_paragraph_after(anchor, text="", style=None, align=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = style
    if align is not None:
        para.alignment = align
    return para


def add_picture_after(anchor, image_path, width_cm=15.0, caption=None):
    para = add_paragraph_after(anchor, "", "Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        cap = add_paragraph_after(para, caption, "Normal",
                                  align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in cap.runs:
            r.italic = True
        return cap
    return para


def add_table_after(anchor, rows, cols):
    tbl = doc.add_table(rows=rows, cols=cols)
    # Bordes manuales (la guia no tiene el estilo "Table Grid")
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)
    anchor._p.addnext(tbl._tbl)
    return tbl


def normalize(text):
    import unicodedata
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().upper()


def find_paragraph(doc, target):
    t = normalize(target)
    for p in doc.paragraphs:
        if t in normalize(p.text):
            return p
    return None


# ------------------------- 1. Copia fresca -------------------------
shutil.copy2(ORIGINAL, OUTPUT)
doc = Document(OUTPUT)

# ------------------------- 2. Portada -------------------------
title_p = find_paragraph(doc, "Evidencia de conocimiento")
if title_p:
    for r in list(title_p.runs):
        r._r.getparent().remove(r._r)
    title_p.add_run(CODIGO)
    anchor = add_paragraph_after(title_p, TITULO_EVIDENCIA, "Normal",
                                 align=WD_ALIGN_PARAGRAPH.CENTER)

fecha_p = find_paragraph(doc, "MES")
if fecha_p:
    for r in list(fecha_p.runs):
        r._r.getparent().remove(r._r)
    fecha_p.add_run(FECHA)

# ------------------------- 3. Introduccion -------------------------
intro_h = find_paragraph(doc, "INTRODUCCION")
intro_text = (
    "En el desarrollo de software, la interfaz gráfica constituye el medio a través "
    "del cual los usuarios interactúan con el sistema, por lo que su diseño debe "
    "responder a los requerimientos funcionales del software y a las necesidades del "
    "usuario final. La presente evidencia tiene como propósito presentar la interfaz "
    "gráfica implementada para el sistema de gestión de lavadero de vehículos "
    "AquaLavado, tomando como referencia el prototipo de navegación elaborado en la "
    "evidencia EV02 y llevándolo a una interfaz funcional desarrollada con tecnologías "
    "web como HTML5, CSS3 y JavaScript."
)
intro_text2 = (
    "Para el diseño de la interfaz se definieron la paleta de colores, la tipografía, "
    "los estilos, los menús de navegación y las reglas de usabilidad que cumplen con "
    "las necesidades del software a desarrollar. Cada una de las vistas implementadas "
    "—inicio de sesión, panel principal, registro de lavados, gestión de clientes, "
    "control de procesos y reportes— se describe a continuación con sus respectivos "
    "pantallazos, evidenciando la funcionalidad de la interfaz gráfica del sistema."
)
if intro_h:
    anchor = add_paragraph_after(intro_h, intro_text, "Normal")
    anchor = add_paragraph_after(anchor, intro_text2, "Normal")

# ------------------------- 4. Desarrollo -------------------------
dev_h = find_paragraph(doc, "DESARROLLO DE LA ACTIVIDAD")

# 4.1 Objetivo
anchor = add_paragraph_after(dev_h, "Objetivo", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Implementar la interfaz gráfica del sistema de gestión de lavadero de vehículos "
    "AquaLavado, definiendo la paleta de colores, los estilos, la usabilidad y los "
    "menús de navegación de acuerdo con los requerimientos del software y las "
    "tecnologías web vistas en los componentes formativos.", "Normal")
for obj in [
    "Definir un sistema de diseño (colores, tipografía y estilos) coherente y acorde con la imagen del negocio.",
    "Construir una interfaz funcional con menús de navegación que permitan recorrer los módulos del sistema.",
    "Aplicar principios de usabilidad para garantizar una experiencia de usuario clara, intuitiva y accesible.",
]:
    anchor = add_paragraph_after(anchor, "• " + obj, "Normal")

# 4.2 Tecnologias
anchor = add_paragraph_after(anchor, "Tecnologías utilizadas", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La interfaz gráfica fue desarrollada con las siguientes tecnologías web: HTML5 "
    "para la estructura de las páginas, CSS3 para el diseño visual (colores, estilos, "
    "posicionamiento y adaptación a diferentes tamaños de pantalla) y JavaScript para "
    "la funcionalidad, la validación de formularios, la navegación entre vistas y la "
    "persistencia de los datos mediante el almacenamiento local del navegador "
    "(localStorage).", "Normal")

# 4.3 Sistema de diseño: paleta de colores
anchor = add_paragraph_after(anchor, "Sistema de diseño: paleta de colores", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Se definió una paleta de colores acorde con la actividad del lavadero, donde el "
    "azul profundo transmite confianza y profesionalismo, y el cian (color del agua) "
    "hace referencia al servicio de lavado. Los colores se aplican de forma "
    "consistente en todos los elementos de la interfaz:", "Normal")
paleta = add_table_after(anchor, 9, 4)
filas = [
    ("Rol", "Color", "Código hexadecimal", "Uso en la interfaz"),
    ("Color primario", "Azul profundo", "#0e3a7d", "Barra lateral, botones y títulos"),
    ("Primario oscuro", "Azul marino", "#0b2a5b", "Degradado de la barra lateral"),
    ("Color de acento", "Cian (agua)", "#22d3ee", "Detalles, resaltados y menú activo"),
    ("Éxito", "Verde", "#10b981", "Lavados finalizados"),
    ("Advertencia", "Ámbar", "#f59e0b", "Lavados en cola de espera"),
    ("Fondo", "Gris claro", "#eef2f7", "Fondo general de la aplicación"),
    ("Superficie", "Blanco", "#ffffff", "Tarjetas y paneles de contenido"),
    ("Texto", "Gris oscuro", "#1e293b", "Textos y encabezados"),
]
for f in filas:
    fila = paleta.rows[filas.index(f)]
    for j, val in enumerate(f):
        fila.cells[j].text = val
# Anclar despues de la tabla para seguir insertando contenido
_p_after_tbl = OxmlElement("w:p")
paleta._tbl.addnext(_p_after_tbl)
anchor = Paragraph(_p_after_tbl, paleta._parent)

# 4.4 Tipografia y usabilidad
anchor = add_paragraph_after(anchor, "Tipografía y usabilidad", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La tipografía utilizada es Segoe UI, una fuente del sistema legible y de estilo "
    "moderno. Los títulos se presentan en tamaño 22px en color azul, los subtítulos "
    "en 16px y el cuerpo del texto entre 13 y 14px, garantizando una lectura cómoda. "
    "En cuanto a usabilidad, la interfaz cuenta con una barra lateral fija que "
    "permite navegar entre los módulos de forma rápida, formularios con validación de "
    "campos obligatorios, mensajes de confirmación al guardar los registros, "
    "etiquetas de color que identifican el estado de cada lavado y un diseño "
    "adaptable (responsivo) a diferentes resoluciones de pantalla.", "Normal")

# 4.5 Menus
anchor = add_paragraph_after(anchor, "Menús de navegación", "Heading 2")
anchor = add_paragraph_after(anchor,
    "La aplicación presenta un menú lateral con los siguientes módulos del sistema: "
    "Panel principal, Registro de lavados, Gestión de clientes, Control de procesos y "
    "Reportes, además de la opción de cerrar sesión. La selección de cada opción "
    "muestra la vista correspondiente sin recargar la página, facilitando el uso del "
    "software.", "Normal")

# 4.6 Pantallas de la interfaz
anchor = add_paragraph_after(anchor, "Pantallas de la interfaz gráfica", "Heading 2")
anchor = add_paragraph_after(anchor,
    "A continuación se presentan los pantallazos de las vistas implementadas, con la "
    "descripción de cada una de las interfaces:", "Normal")

interfaces = [
    ("1_login.png", "Interfaz 1. Inicio de sesión (Login)",
     "Vista de acceso al sistema. Valida las credenciales del usuario; al ingresar "
     "las credenciales correctas (admin / admin123) inicia la sesión y redirige al "
     "panel principal. Incluye mensaje de error cuando los datos son incorrectos y "
     "mantiene la sesión activa durante el uso de la aplicación."),
    ("2_dashboard.png", "Interfaz 2. Panel principal (Dashboard)",
     "Resumen general de la operación con tarjetas de indicadores: lavados del día, "
     "ingresos generados, clientes registrados y lavados en proceso. Incluye un "
     "gráfico de barras con los lavados por tipo de servicio, calculado a partir de "
     "los datos registrados."),
    ("3_lavados.png", "Interfaz 3. Registro de lavados",
     "Permite registrar un nuevo lavado seleccionando el cliente, el vehículo, el "
     "tipo de servicio (cada uno con su valor) y el operario asignado. Al guardar, el "
     "lavado aparece en la tabla de lavados recientes y los datos se almacenan en el "
     "navegador mediante localStorage."),
    ("4_clientes.png", "Interfaz 4. Gestión de clientes",
     "Muestra el listado de clientes del lavadero con buscador en tiempo real por "
     "nombre, cédula o vehículo. Incluye un formulario para registrar nuevos clientes, "
     "validando los campos obligatorios y mostrando un mensaje de confirmación."),
    ("5_procesos.png", "Interfaz 5. Control de procesos",
     "Tablero kanban con las etapas del proceso de lavado: en cola, en proceso y "
     "finalizados. Cada tarjeta muestra el vehículo, el tipo de servicio y el "
     "operario, y el botón \u201cAvanzar\u201d permite cambiar de estado cada lavado, "
     "actualizando el tablero en tiempo real."),
    ("6_reportes.png", "Interfaz 6. Reportes y estadísticas",
     "Consolida los indicadores de la operación: ingresos del periodo, total de "
     "lavados, ticket promedio y clientes. Presenta un gráfico de ingresos por tipo "
     "de lavado y una tabla con el comportamiento de los lavados por día de la "
     "semana, facilitando la toma de decisiones."),
]

for i, (img, titulo, desc) in enumerate(interfaces, start=1):
    anchor = add_paragraph_after(anchor, titulo, "Heading 2")
    img_path = CAPTURAS / img
    if img_path.exists():
        nombre_interfaz = titulo.split(". ", 1)[1].lower()
        anchor = add_picture_after(
            anchor, img_path, width_cm=14.5,
            caption=f"Figura {i}. Interfaz {nombre_interfaz} — implementación real.")
    else:
        anchor = add_paragraph_after(anchor, f"[Pantallazo pendiente: {img}]", "Normal")
    anchor = add_paragraph_after(anchor, desc, "Normal")

# ------------------------- 5. Conclusion -------------------------
concl_h = find_paragraph(doc, "CONCLUSION")
concl_text = (
    "La elaboración de esta evidencia permitió transformar el prototipo de navegación "
    "en una interfaz gráfica funcional, aplicando los componentes y tecnologías vistos "
    "en los componentes formativos. A través de este ejercicio se comprendió la "
    "importancia de definir un sistema de diseño coherente —colores, tipografía, "
    "estilos y menús— que no solo cumpla con los requerimientos del software, sino que "
    "también garantice una experiencia de usuario agradable e intuitiva."
)
concl_text2 = (
    "Se fortalecieron competencias en el desarrollo de interfaces con HTML5, CSS3 y "
    "JavaScript, así como en la aplicación de principios de usabilidad y en la "
    "implementación de funcionalidades como la validación de formularios, la "
    "navegación entre vistas y la persistencia de datos. Se concluye que la interfaz "
    "gráfica del sistema AquaLavado cumple con las necesidades del software a "
    "desarrollar y constituye la base para la construcción del aplicativo completo."
)
if concl_h:
    anchor = add_paragraph_after(concl_h, concl_text, "Normal")
    anchor = add_paragraph_after(anchor, concl_text2, "Normal")

doc.save(OUTPUT)
print(f"OK: documento generado en {OUTPUT}")
missing = [img for img, _, _ in interfaces if not (CAPTURAS / img).exists()]
if missing:
    print("Faltan capturas:", ", ".join(missing))
else:
    print("Todas las capturas fueron insertadas.")

