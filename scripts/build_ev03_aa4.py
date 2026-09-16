# -*- coding: utf-8 -*-
"""
Construye la evidencia GA6-220501096-AA4-EV03 (diseño front-end - Lavadero).

Siempre parte de la guia original (1GUIA_PARA_SENA .docx) para que sea
idempotente: se puede ejecutar varias veces sin duplicar contenido.
Inserta capturas de la interfaz real (interfaz/capturas) si existen.
"""
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ORIGINAL = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\1GUIA_PARA_SENA .docx"
OUTPUT = r"C:\Users\dylan\Documents\OpenCode\Act-OpenCode\GA6-220501096-AA4-EV03.docx"
CAPTURAS = Path(r"C:\Users\dylan\Documents\OpenCode\ADSO\proyectos\lavadero\interfaz\capturas")

CODIGO = "GA6-220501096-AA4-EV03"
TITULO_EVIDENCIA = "Diseño front-end que cumpla con los requerimientos del proyecto"
FECHA = "Agosto 2026"

# ------------------------- Utilidades -------------------------

def add_paragraph_after(anchor, text="", style=None, align=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = style
    if align is not None:
        para.alignment = align
    return para


def add_picture_after(anchor, image_path, width_cm=15.5, caption=None):
    para = add_paragraph_after(anchor, "", "Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        cap = add_paragraph_after(para, caption, "Normal",
                                  align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in cap.runs:
            r.italic = True
        return cap
    return para


def add_table_after(anchor, rows, cols, width_cols=None):
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)
    if width_cols:
        for row in tbl.rows:
            for idx, w in enumerate(width_cols):
                row.cells[idx].width = Cm(w)
    anchor._p.addnext(tbl._tbl)
    return tbl


def anchor_after_table(tbl):
    new_p = OxmlElement("w:p")
    tbl._tbl.addnext(new_p)
    return Paragraph(new_p, tbl._parent)


def fill_table(tbl, data):
    for r, fila in enumerate(data):
        for c, val in enumerate(fila):
            cell = tbl.rows[r].cells[c]
            cell.text = val
            if r == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def add_code_paragraph(anchor, code, style="Normal", align=None):
    para = add_paragraph_after(anchor, "", style, align)
    run = para.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Cm(0.40)
    return para


def add_code_block(anchor, code):
    for linea in code.split("\n"):
        anchor = add_code_paragraph(anchor, linea)
    return anchor


def normalize(text):
    import unicodedata
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().upper()


def find_paragraph(doc, target):
    t = normalize(target)
    for p in doc.paragraphs:
        if t in normalize(p.text):
            return p
    return None


# ------------------------- 1. Copia fresca desde la guia original -------------------------
shutil.copy2(ORIGINAL, OUTPUT)
doc = Document(OUTPUT)

# ------------------------- 2. Portada -------------------------
title_p = find_paragraph(doc, "Evidencia de conocimiento")
if title_p:
    for r in list(title_p.runs):
        r._r.getparent().remove(r._r)
    title_p.add_run(CODIGO)
    anchor = add_paragraph_after(title_p, TITULO_EVIDENCIA, "Normal",
                                 align=WD_ALIGN_PARAGRAPH.CENTER)

fecha_p = find_paragraph(doc, "MES")
if fecha_p:
    for r in list(fecha_p.runs):
        r._r.getparent().remove(r._r)
    fecha_p.add_run(FECHA)

# ------------------------- 3. Introduccion -------------------------
intro_h = find_paragraph(doc, "INTRODUCCION")
intro_text = (
    "El diseño front-end es la materialización visual e interactiva de los "
    "requerimientos del software, y su calidad depende de la correcta aplicación de "
    "los conceptos de usabilidad, de los componentes de diseño y de las tecnologías "
    "web HTML, CSS y JavaScript. La presente evidencia tiene como propósito presentar "
    "el diseño front-end de la aplicación web del sistema de gestión de lavadero de "
    "vehículos AquaLavado, cumpliendo con los requisitos del proyecto y aplicando lo "
    "visto en el componente formativo."
)
intro_text2 = (
    "Para dar cumplimiento a la actividad se tomó como base el prototipo de "
    "navegación y la interfaz gráfica desarrollados en las evidencias anteriores, "
    "aplicando los conceptos de usabilidad y diseño allí propuestos: una paleta de "
    "colores coherente, tipografía legible, menús de navegación claros, formularios "
    "validados, estados visuales para cada proceso y un diseño responsivo. El "
    "front-end fue implementado con HTML para la estructura, CSS para los estilos y "
    "JavaScript para la funcionalidad, tal como se describe a continuación."
)
if intro_h:
    anchor = add_paragraph_after(intro_h, intro_text, "Normal")
    anchor = add_paragraph_after(anchor, intro_text2, "Normal")

# ------------------------- 4. Desarrollo de la actividad -------------------------
dev_h = find_paragraph(doc, "DESARROLLO DE LA ACTIVIDAD")

# 4.1 Objetivo
anchor = add_paragraph_after(dev_h, "Objetivo", "Heading 2")
anchor = add_paragraph_after(anchor,
    "Realizar el diseño front-end de la aplicación web del sistema de gestión de "
    "lavadero de vehículos AquaLavado, cumpliendo con los requisitos del proyecto y "
    "aplicando los conceptos de usabilidad y diseño propuestos en las evidencias "
    "anteriores, mediante el uso de HTML, CSS y JavaScript.", "Normal")
for obj in [
    "Aplicar los conceptos de usabilidad y diseño establecidos en las evidencias anteriores (paleta de colores, tipografía, menús y estados visuales).",
    "Implementar el diseño front-end con HTML para la estructura, CSS para los estilos y JavaScript para la funcionalidad.",
    "Cumplir con los requerimientos funcionales del software en cada una de las vistas del sistema.",
    "Entregar la carpeta con todos los archivos del diseño front-end comprimida en formato ZIP.",
]:
    anchor = add_paragraph_after(anchor, "• " + obj, "Normal")

# 4.2 Requerimientos del proyecto
anchor = add_paragraph_after(anchor, "Requerimientos del proyecto", "Heading 2")
anchor = add_paragraph_after(anchor,
    "El sistema AquaLavado debe permitir la gestión de la operación de un lavadero de "
    "vehículos. Los requerimientos funcionales que orientan el diseño front-end son:", "Normal")
for req in [
    "Inicio de sesión: permitir el acceso al sistema mediante usuario y contraseña.",
    "Panel principal: mostrar indicadores de la operación (lavados, ingresos, clientes y procesos).",
    "Registro de lavados: capturar el cliente, el vehículo, el tipo de servicio y el operario.",
    "Gestión de clientes: registrar y buscar clientes con sus datos y vehículos.",
    "Control de procesos: dar seguimiento al estado de cada lavado (en cola, en proceso, finalizado).",
    "Reportes: presentar estadísticas e indicadores de ingresos y lavados.",
]:
    anchor = add_paragraph_after(anchor, "• " + req, "Normal")

# 4.3 Conceptos de usabilidad y diseño aplicados
anchor = add_paragraph_after(anchor, "Conceptos de usabilidad y diseño aplicados", "Heading 2")
anchor = add_paragraph_after(anchor,
    "El diseño front-end aplica los conceptos de usabilidad y diseño propuestos en las "
    "evidencias anteriores, garantizando una experiencia de usuario clara e "
    "intuitiva:", "Normal")
for uso in [
    "Paleta de colores: azul profundo (#0e3a7d) y cian (#22d3ee) que transmiten confianza y hacen referencia al agua del lavado.",
    "Tipografía legible (Segoe UI) con tamaños jerárquicos para títulos, subtítulos y cuerpo del texto.",
    "Menú de navegación lateral fijo que permite acceder a los módulos del sistema de forma rápida.",
    "Formularios con validación de campos y mensajes de confirmación al guardar los datos.",
    "Etiquetas de estado por color que identifican cada etapa del proceso de lavado.",
    "Tablas con encabezados, filas alternadas y efecto hover para facilitar la lectura.",
    "Diseño responsivo que se adapta a diferentes tamaños de pantalla (media queries).",
    "Retroalimentación visual (hover, sombras y transiciones) en los elementos interactivos.",
]:
    anchor = add_paragraph_after(anchor, "• " + uso, "Normal")

# 4.4 Tecnologias y estructura
anchor = add_paragraph_after(anchor, "Tecnologías y estructura de archivos", "Heading 2")
anchor = add_paragraph_after(anchor,
    "El diseño front-end se implementa con las tecnologías HTML, CSS y JavaScript, "
    "organizadas en la siguiente estructura de archivos:", "Normal")
anchor = add_code_block(anchor,
    "interfaz/\n"
    "├── index.html          → Página de inicio de sesión (HTML)\n"
    "├── app.html            → Aplicación principal (HTML)\n"
    "├── css/\n"
    "│   └── estilo.css      → Hoja de estilos CSS\n"
    "└── js/\n"
    "    ├── datos.js        → Datos iniciales y configuración\n"
    "    └── app.js          → Lógica y funcionalidad (JavaScript)")

# 4.5 Implementacion con HTML
anchor = add_paragraph_after(anchor, "Implementación con HTML", "Heading 2")
anchor = add_paragraph_after(anchor,
    "HTML se utiliza para definir la estructura de la interfaz: etiquetas semánticas "
    "(<header>, <nav>, <main>, <section>), formularios (<form>, <input>, <select>, "
    "<button>), tablas (<table>, <thead>, <tbody>) y contenedores (<div>) para las "
    "tarjetas y el tablero kanban. A continuación se muestra un fragmento del "
    "documento app.html con la estructura de la aplicación:", "Normal")
anchor = add_code_block(anchor,
    "<body>\n"
    "  <aside class=\"barra-lateral\">\n"
    "    <nav>\n"
    "      <a href=\"#panel\">Panel principal</a>\n"
    "      <a href=\"#lavados\">Registro de lavados</a>\n"
    "      <a href=\"#clientes\">Gestión de clientes</a>\n"
    "      <a href=\"#procesos\">Control de procesos</a>\n"
    "      <a href=\"#reportes\">Reportes</a>\n"
    "    </nav>\n"
    "  </aside>\n"
    "  <main>\n"
    "    <section id=\"panel\">...</section>\n"
    "    <section id=\"lavados\">...</section>\n"
    "    <section id=\"clientes\">...</section>\n"
    "    <section id=\"procesos\">...</section>\n"
    "    <section id=\"reportes\">...</section>\n"
    "  </main>\n"
    "</body>")

# 4.6 Implementacion con CSS
anchor = add_paragraph_after(anchor, "Implementación con CSS", "Heading 2")
anchor = add_paragraph_after(anchor,
    "CSS define la presentación visual del front-end: variables de color, el modelo "
    "de caja, flexbox y grid para la maquetación, y media queries para el diseño "
    "responsivo. A continuación se muestra un fragmento de la hoja estilo.css con la "
    "paleta de colores y la barra lateral:", "Normal")
anchor = add_code_block(anchor,
    ":root {\n"
    "  --color-primario: #0e3a7d;\n"
    "  --color-acento: #22d3ee;\n"
    "  --color-fondo: #eef2f7;\n"
    "}\n"
    "\n"
    ".barra-lateral {\n"
    "  width: 230px;\n"
    "  background: linear-gradient(135deg, #0b2a5b, #0e3a7d);\n"
    "  color: #ffffff;\n"
    "  position: fixed;\n"
    "  height: 100vh;\n"
    "}\n"
    "\n"
    "@media (max-width: 768px) {\n"
    "  .barra-lateral { width: 60px; }\n"
    "}")

# 4.7 Implementacion con JavaScript
anchor = add_paragraph_after(anchor, "Implementación con JavaScript", "Heading 2")
anchor = add_paragraph_after(anchor,
    "JavaScript brinda la funcionalidad del front-end: validación de la sesión, "
    "navegación entre vistas, gestión de formularios, renderizado dinámico de tablas "
    "y del tablero kanban, búsqueda en tiempo real y persistencia de datos en "
    "localStorage. A continuación se muestra un fragmento del archivo app.js con la "
    "navegación entre secciones y la función para avanzar el estado de un lavado:", "Normal")
anchor = add_code_block(anchor,
    "// Navegación entre secciones\n"
    "function mostrarSeccion(id) {\n"
    "  document.querySelectorAll('main section')\n"
    "    .forEach(s => s.style.display = 'none');\n"
    "  document.getElementById(id).style.display = 'block';\n"
    "}\n"
    "\n"
    "// Avanzar estado de un lavado en el kanban\n"
    "function avanzarEstado(id) {\n"
    "  const lavado = lavados.find(l => l.id === id);\n"
    "  const orden = ['cola', 'proceso', 'finalizado'];\n"
    "  const actual = orden.indexOf(lavado.estado);\n"
    "  lavado.estado = orden[actual + 1];\n"
    "  renderKanban();\n"
    "  guardarDatos();\n"
    "}")

# 4.8 Vista de las pantallas del front-end
anchor = add_paragraph_after(anchor, "Pantallas del diseño front-end", "Heading 2")
anchor = add_paragraph_after(anchor,
    "A continuación se presentan las pantallas del diseño front-end implementado, "
    "donde se evidencia el cumplimiento de los requerimientos del proyecto y la "
    "aplicación de los conceptos de usabilidad y diseño:", "Normal")

interfaces = [
    ("1_login.png", "Pantalla 1. Inicio de sesión",
     "Cumple el requerimiento de acceso al sistema. Aplica el formulario con "
     "validación de credenciales, mensaje de error en datos incorrectos y los "
     "colores institucionales de la paleta definida."),
    ("2_dashboard.png", "Pantalla 2. Panel principal",
     "Muestra los indicadores de la operación en tarjetas con la paleta de colores "
     "del sistema y el menú lateral de navegación. Incluye un gráfico de lavados por "
     "tipo de servicio."),
    ("3_lavados.png", "Pantalla 3. Registro de lavados",
     "Formulario con selección de cliente, vehículo, tipo de servicio y operario; la "
     "tabla de lavados recientes muestra los registros con su estado mediante "
     "etiquetas de color."),
    ("4_clientes.png", "Pantalla 4. Gestión de clientes",
     "Listado de clientes con buscador en tiempo real y formulario de registro con "
     "validación de campos obligatorios y mensaje de confirmación."),
    ("5_procesos.png", "Pantalla 5. Control de procesos",
     "Tablero kanban con las etapas del lavado (en cola, en proceso, finalizado). El "
     "botón \"Avanzar\" permite cambiar de estado cada lavado, actualizando el "
     "tablero en tiempo real."),
    ("6_reportes.png", "Pantalla 6. Reportes y estadísticas",
     "Consolida los indicadores de ingresos, total de lavados y ticket promedio, con "
     "gráficos y tablas que facilitan la toma de decisiones."),
]

for i, (img, titulo, desc) in enumerate(interfaces, start=1):
    anchor = add_paragraph_after(anchor, titulo, "Heading 2")
    img_path = CAPTURAS / img
    if img_path.exists():
        anchor = add_picture_after(anchor, img_path, width_cm=15.0)
    else:
        anchor = add_paragraph_after(anchor, f"[Pantallazo pendiente: {img}]", "Normal")
    anchor = add_paragraph_after(anchor, desc, "Normal")

# 4.9 Archivos entregados
anchor = add_paragraph_after(anchor, "Archivos entregados", "Heading 2")
anchor = add_paragraph_after(anchor,
    "De acuerdo con los lineamientos de la evidencia, se entrega una carpeta con "
    "todos los archivos del diseño front-end (HTML, CSS y JavaScript), comprimida en "
    "formato ZIP. Los archivos de la carpeta son:", "Normal")
for f in [
    "index.html — página de inicio de sesión.",
    "app.html — aplicación principal con los módulos del sistema.",
    "css/estilo.css — hoja de estilos CSS del front-end.",
    "js/datos.js — datos iniciales y configuración.",
    "js/app.js — lógica y funcionalidad JavaScript.",
]:
    anchor = add_paragraph_after(anchor, "• " + f, "Normal")

# ------------------------- 5. Conclusion -------------------------
concl_h = find_paragraph(doc, "CONCLUSION")
concl_text = (
    "El desarrollo de la presente evidencia permitió materializar el diseño front-end "
    "de la aplicación web AquaLavado, cumpliendo con los requerimientos del proyecto "
    "y aplicando los conceptos de usabilidad y diseño propuestos en las evidencias "
    "anteriores. Se implementó una interfaz funcional con HTML para la estructura, "
    "CSS para los estilos y JavaScript para la funcionalidad, demostrando la "
    "aplicación práctica de los conocimientos del componente formativo."
)
concl_text2 = (
    "Así mismo, se fortalecieron competencias en el desarrollo de front-end, la "
    "usabilidad, el diseño responsivo y la programación de la interfaz, integrando "
    "los componentes definidos en el documento de la evidencia AA4-EV02. Se concluye "
    "que el diseño front-end cumple con los requisitos del software AquaLavado y "
    "queda listo para su integración con el backend de la aplicación."
)
if concl_h:
    anchor = add_paragraph_after(concl_h, concl_text, "Normal")
    anchor = add_paragraph_after(anchor, concl_text2, "Normal")

doc.save(OUTPUT)
print(f"OK: documento generado en {OUTPUT}")
missing = [img for img, _, _ in interfaces if not (CAPTURAS / img).exists()]
if missing:
    print("Faltan capturas:", ", ".join(missing))
else:
    print("Todas las capturas fueron insertadas.")

